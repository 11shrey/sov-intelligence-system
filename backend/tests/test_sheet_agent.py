"""Tests for Agent 1 (Sheet Intelligence Agent)."""

import json
from pathlib import Path
from openpyxl import Workbook
import pytest

from app.agents.sheet_agent import SheetIntelligenceAgent
from app.orchestration.state import SOVProcessingState, FileInfo, JobStatus


@pytest.fixture
def agent():
    return SheetIntelligenceAgent()


def create_state_for_file(file_path: Path) -> SOVProcessingState:
    file_info = FileInfo(
        filename=file_path.name,
        file_path=str(file_path),
        file_size_bytes=file_path.stat().st_size if file_path.exists() else 100,
    )
    return SOVProcessingState(job_id="test-sheet-job", file_info=file_info)


def test_multi_sheet_selection(agent, tmp_path):
    """Verify Agent 1 selects the property schedule over Instructions and Summary sheets."""
    wb = Workbook()

    # Sheet 1: Instructions
    ws_inst = wb.active
    ws_inst.title = "Instructions"
    ws_inst.append(["Please read instructions before filling the values"])
    ws_inst.append(["Contact underwriter at support@carrier.com"])

    # Sheet 2: Summary
    ws_summary = wb.create_sheet(title="Executive Summary")
    ws_summary.append(["Total Portfolio TIV", 50000000])
    ws_summary.append(["Total Locations Count", 25])

    # Sheet 3: Real SOV Schedule
    ws_sov = wb.create_sheet(title="Location Schedule")
    ws_sov.append(["Loc #", "Address", "City", "State", "Zip", "Building Value", "Occupancy", "Construction"])
    ws_sov.append(["LOC-1", "100 Main St", "Chicago", "IL", "60601", 2000000, "Office", "Masonry"])
    ws_sov.append(["LOC-2", "200 State St", "Boston", "MA", "02109", 3500000, "Retail", "Frame"])
    ws_sov.append(["LOC-3", "300 Market St", "Philadelphia", "PA", "19106", 1800000, "Warehouse", "Steel"])

    wb_path = tmp_path / "multi_sheet_portfolio.xlsx"
    wb.save(wb_path)

    state = create_state_for_file(wb_path)
    updated_state = agent.run(state)

    assert updated_state.status == JobStatus.SHEET_ANALYZED
    assert updated_state.selected_sheet == "Location Schedule"
    assert updated_state.header_row == 0
    assert len(updated_state.sheet_analysis) == 3

    # Check that candidate analysis has high confidence for Location Schedule
    sov_analysis = next(a for a in updated_state.sheet_analysis if a.sheet_name == "Location Schedule")
    assert sov_analysis.is_candidate is True
    assert sov_analysis.confidence >= 0.70


def test_header_on_row_5_with_blank_leading_rows(agent, tmp_path):
    """Verify header detection when title/banner and blank rows precede the table (header on 0-indexed row 4)."""
    # 1. Test in XLSX
    wb = Workbook()
    ws = wb.active
    ws.title = "Property Data"
    ws.append([])  # Row 0: Blank
    ws.append(["CONFIDENTIAL STATEMENT OF VALUES"])  # Row 1: Banner/Title
    ws.append(["As of October 2026 - Broker Submission"])  # Row 2: Metadata
    ws.append([])  # Row 3: Blank
    ws.append(["Reference", "Property Address", "City", "State", "Zip Code", "Bldg Value ($)", "Occupancy"])  # Row 4: Header
    ws.append(["PROP-01", "123 Industrial Pkwy", "Dallas", "TX", "75001", 4500000, "Manufacturing"])
    ws.append(["PROP-02", "456 Commerce Ave", "Houston", "TX", "77002", 6200000, "Office"])

    xlsx_path = tmp_path / "banner_header_row_4.xlsx"
    wb.save(xlsx_path)

    state_xlsx = create_state_for_file(xlsx_path)
    res_xlsx = agent.run(state_xlsx)
    assert res_xlsx.selected_sheet == "Property Data"
    assert res_xlsx.header_row == 4

    # 2. Test in CSV (guards against skip_blank_lines)
    csv_content = (
        "\n"
        "CONFIDENTIAL STATEMENT OF VALUES\n"
        "As of October 2026 - Broker Submission\n"
        "\n"
        "Reference,Property Address,City,State,Zip Code,Bldg Value ($),Occupancy\n"
        "PROP-01,123 Industrial Pkwy,Dallas,TX,75001,4500000,Manufacturing\n"
        "PROP-02,456 Commerce Ave,Houston,TX,77002,6200000,Office\n"
    )
    csv_path = tmp_path / "banner_header_row_4.csv"
    csv_path.write_text(csv_content, encoding="utf-8")

    state_csv = create_state_for_file(csv_path)
    res_csv = agent.run(state_csv)
    assert res_csv.selected_sheet == "CSV"
    assert res_csv.header_row == 4


def test_coverage_schedule_not_rejected_by_cover(agent, tmp_path):
    """Verify sheet named 'Coverage Schedule' is NOT rejected by 'cover'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage Schedule"
    ws.append(["Loc ID", "Street Address", "City", "State", "Zip", "Building Limit", "Contents Limit"])
    ws.append(["1", "500 Lake Shore Dr", "Chicago", "IL", "60611", 12000000, 3000000])

    wb_path = tmp_path / "coverage_schedule.xlsx"
    wb.save(wb_path)

    state = create_state_for_file(wb_path)
    res = agent.run(state)

    assert res.selected_sheet == "Coverage Schedule"
    assert res.header_row == 0
    analysis = res.sheet_analysis[0]
    assert analysis.is_candidate is True
    assert "matches non-data sheet" not in analysis.reasoning


def test_statement_header_does_not_match_state(agent, tmp_path):
    """Verify column named 'Statement' does NOT match SOV concept 'State'."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Legal Disclaimer"
    ws.append(["Statement", "Notes", "Paragraph", "Author"])
    ws.append(["This is an informative statement.", "No liability assumed.", "1", "Legal Team"])

    wb_path = tmp_path / "legal.xlsx"
    wb.save(wb_path)

    state = create_state_for_file(wb_path)
    res = agent.run(state)

    assert res.selected_sheet is None
    assert res.header_row is None
    assert len(res.errors) > 0


def test_no_valid_sov_sheet_graceful_handling(agent, tmp_path):
    """Verify graceful result with no exception when workbook contains no SOV data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Notes & Contact"
    ws.append(["Meeting Notes", "Date", "Attendees"])
    ws.append(["Discussed renewal terms", "2026-10-01", "Alex, Jordan"])

    wb_path = tmp_path / "notes_only.xlsx"
    wb.save(wb_path)

    state = create_state_for_file(wb_path)
    res = agent.run(state)

    # Must not crash
    assert res.status == JobStatus.SHEET_ANALYZED
    assert res.selected_sheet is None
    assert res.header_row is None
    assert any("No valid SOV candidate" in err for err in res.errors)


def test_uniform_detection_across_xlsx_csv_json(agent, tmp_path):
    """Verify messy xlsx, csv, and json of same property data produce uniform detection."""
    records = [
        {"Reference": "L-1", "Address": "100 King St", "City": "Toronto", "State": "ON", "Zip": "M5X 1A9", "Building Value": 15000000},
        {"Reference": "L-2", "Address": "200 Bay St", "City": "Toronto", "State": "ON", "Zip": "M5J 2J2", "Building Value": 22000000},
    ]

    # 1. JSON
    json_path = tmp_path / "data.json"
    json_path.write_text(json.dumps(records), encoding="utf-8")
    res_json = agent.run(create_state_for_file(json_path))
    assert res_json.selected_sheet == "JSON"
    assert res_json.header_row == 0

    # 2. CSV
    csv_path = tmp_path / "data.csv"
    csv_path.write_text("Reference,Address,City,State,Zip,Building Value\nL-1,100 King St,Toronto,ON,M5X 1A9,15000000\nL-2,200 Bay St,Toronto,ON,M5J 2J2,22000000\n", encoding="utf-8")
    res_csv = agent.run(create_state_for_file(csv_path))
    assert res_csv.selected_sheet == "CSV"
    assert res_csv.header_row == 0

    # 3. XLSX
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.append(["Reference", "Address", "City", "State", "Zip", "Building Value"])
    ws.append(["L-1", "100 King St", "Toronto", "ON", "M5X 1A9", 15000000])
    ws.append(["L-2", "200 Bay St", "Toronto", "ON", "M5J 2J2", 22000000])
    xlsx_path = tmp_path / "data.xlsx"
    wb.save(xlsx_path)
    res_xlsx = agent.run(create_state_for_file(xlsx_path))
    assert res_xlsx.selected_sheet == "Schedule"
    assert res_xlsx.header_row == 0
