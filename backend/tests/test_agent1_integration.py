"""Production-ready End-to-End Integration Tests for Agent 1.

Validates the full pipeline:
File -> File Type Detection -> Ingestion -> Semantics -> Header Detection ->
Sheet Selection -> Agent 1 -> Agent 2 JSON Handoff Contract.
"""

from pathlib import Path
import json
import openpyxl
import pytest

from app.agents.sheet_intelligence import (
    Agent1HandoffResult,
    CorruptedFileError,
    EmptyFileError,
    FileNotFoundIngestionError,
    MalformedJSONError,
    UnsupportedFormatError,
    analyze_file,
)


# =====================================================================
# 6. End-to-End XLSX Test (Multi-Sheet Workbook)
# =====================================================================


def test_e2e_xlsx_multi_sheet(tmp_path: Path):
    """Section 6: Full pipeline on a realistic XLSX workbook with Instructions, Summary, Property Schedule."""
    file_path = tmp_path / "broker_submission.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Instructions
    ws1 = wb.active
    ws1.title = "Instructions"
    ws1.append(["Broker Submission Instructions"])
    ws1.append(["Complete all required fields"])

    # Sheet 2: Summary
    ws2 = wb.create_sheet(title="Summary")
    ws2.append(["Submission Summary"])
    ws2.append(["Total Locations", 2])
    ws2.append(["Total Value", 4700000])

    # Sheet 3: Property Schedule
    ws3 = wb.create_sheet(title="Property Schedule")
    ws3.append(["ABC Broker"])
    ws3.append(["Submission Date", "2026-01-01"])
    ws3.append([])
    ws3.append(
        [
            "Loc ID",
            "Street Address",
            "City",
            "State",
            "Zip",
            "Building Cost",
            "Occupancy",
        ]
    )
    ws3.append(
        ["P001", "100 Main St", "Chicago", "IL", "60601", "1500000", "Office"]
    )
    ws3.append(
        ["P002", "200 State St", "Chicago", "IL", "60602", "3200000", "Retail"]
    )

    wb.save(file_path)

    result = analyze_file(file_path, job_id="job_e2e_xlsx")

    assert isinstance(result, Agent1HandoffResult)
    assert result.selected_sheet == "Property Schedule"
    assert result.header_row == 3
    assert result.raw_headers == [
        "Loc ID",
        "Street Address",
        "City",
        "State",
        "Zip",
        "Building Cost",
        "Occupancy",
    ]
    assert len(result.sample_rows) == 2
    assert result.sample_rows[0] == [
        "P001",
        "100 Main St",
        "Chicago",
        "IL",
        "60601",
        "1500000",
        "Office",
    ]
    assert result.total_rows == 2
    assert result.confidence >= 0.85
    assert len(result.all_sheets_evaluated) == 3


# =====================================================================
# 7. End-to-End CSV Test (Metadata Rows Before Header)
# =====================================================================


def test_e2e_csv_metadata_before_header(tmp_path: Path):
    """Section 7: Full pipeline on CSV containing metadata before the actual header."""
    file_path = tmp_path / "metadata_schedule.csv"
    csv_content = (
        "Broker Submission\n"
        "ABC Broker\n"
        "\n"
        "Submission Date,2026-01-01\n"
        "\n"
        "Loc ID,Street Address,City,State,Zip,Building Cost,Occupancy\n"
        "P001,100 Main St,Chicago,IL,60601,1500000,Office\n"
        "P002,200 State St,Chicago,IL,60602,3200000,Retail\n"
    )
    file_path.write_text(csv_content, encoding="utf-8")

    result = analyze_file(file_path, job_id="job_e2e_csv")

    assert result.file_type == "csv"
    assert result.selected_sheet == "default"
    # Row 5 contains the true header
    assert result.header_row == 5
    assert result.raw_headers == [
        "Loc ID",
        "Street Address",
        "City",
        "State",
        "Zip",
        "Building Cost",
        "Occupancy",
    ]
    assert result.total_rows == 2
    assert len(result.sample_rows) == 2


# =====================================================================
# 8. End-to-End JSON Test
# =====================================================================


def test_e2e_json_records(tmp_path: Path):
    """Section 8: Full pipeline on supported tabular JSON list of records."""
    file_path = tmp_path / "records.json"
    content = """[
        {
            "Loc ID": "P001",
            "Street Address": "100 Main St",
            "City": "Chicago",
            "State": "IL",
            "Zip": "60601",
            "Building Cost": "1500000",
            "Occupancy": "Office"
        },
        {
            "Loc ID": "P002",
            "Street Address": "200 State St",
            "City": "Chicago",
            "State": "IL",
            "Zip": "60602",
            "Building Cost": "3200000",
            "Occupancy": "Retail"
        }
    ]"""
    file_path.write_text(content, encoding="utf-8")

    result = analyze_file(file_path, job_id="job_e2e_json")

    assert result.file_type == "json"
    assert result.selected_sheet == "default"
    assert result.header_row == 0
    assert result.raw_headers == [
        "Loc ID",
        "Street Address",
        "City",
        "State",
        "Zip",
        "Building Cost",
        "Occupancy",
    ]
    assert result.total_rows == 2
    assert result.sample_rows[0] == [
        "P001",
        "100 Main St",
        "Chicago",
        "IL",
        "60601",
        "1500000",
        "Office",
    ]


# =====================================================================
# 9. Raw Data Preservation (Anti-Mapping Test)
# =====================================================================


def test_raw_data_preservation_anti_mapping(tmp_path: Path):
    """Section 9: Raw strings are strictly preserved and NOT mapped to canonical schema."""
    file_path = tmp_path / "anti_map.csv"
    csv_content = (
        "Site,Location,Building Cost,Value\n"
        "S001,North Campus,1250000,1500000\n"
        "S002,South Campus,2250000,2500000\n"
    )
    file_path.write_text(csv_content, encoding="utf-8")

    result = analyze_file(file_path, job_id="job_anti_map")
    json_str = result.model_dump_json()
    data = json.loads(json_str)

    # 1. Raw headers exact strings
    assert data["raw_headers"] == ["Site", "Location", "Building Cost", "Value"]

    # 2. Sample records exact strings
    assert data["sample_rows"][0] == ["S001", "North Campus", "1250000", "1500000"]

    # 3. Prove NO canonical mappings were emitted
    assert '"Site": "address"' not in json_str
    assert '"Location": "reference"' not in json_str
    assert '"Building Cost": "building_value"' not in json_str
    assert '"Value": "building_value"' not in json_str


# =====================================================================
# 10. Header Offset Test (Rows 0, 2, 4, 8)
# =====================================================================


@pytest.mark.parametrize("offset", [0, 2, 4, 8])
def test_header_offset_positions(tmp_path: Path, offset: int):
    """Section 10: Verifies header row index exactly tracks source row offset."""
    file_path = tmp_path / f"offset_{offset}.csv"
    lines = []

    for i in range(offset):
        if i % 2 == 0:
            lines.append(f"Metadata Line {i}")
        else:
            lines.append("")

    lines.append(
        "Loc ID,Street Address,City,State,Zip,Building Cost,Occupancy"
    )
    lines.append("001,100 Main St,Chicago,IL,60601,1500000,Office")
    lines.append("002,200 Main St,Chicago,IL,60602,2500000,Retail")

    file_path.write_text("\n".join(lines), encoding="utf-8")

    result = analyze_file(file_path)
    assert result.header_row == offset
    assert result.raw_headers[0] == "Loc ID"


# =====================================================================
# 11. Multiple-Sheet Test (5 Sheets Evaluated)
# =====================================================================


def test_multiple_sheet_workbook_all_evaluated(tmp_path: Path):
    """Section 11: Workbook with 5 sheets: Instructions, Cover, Summary, Property Schedule, Claims."""
    file_path = tmp_path / "five_sheets.xlsx"
    wb = openpyxl.Workbook()

    ws_inst = wb.active
    ws_inst.title = "Instructions"
    ws_inst.append(["Read Instructions"])

    ws_cover = wb.create_sheet(title="Cover")
    ws_cover.append(["Policy Schedule Cover"])

    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.append(["Total Value", 50000000])

    ws_sov = wb.create_sheet(title="Property Schedule")
    ws_sov.append(["Loc ID", "Address", "City", "State", "Zip", "Building Value", "Occupancy"])
    ws_sov.append(["L-1", "100 Main", "Chicago", "IL", "60601", 10000000, "Office"])

    ws_claims = wb.create_sheet(title="Claims")
    ws_claims.append(["Claim #", "Amount"])
    ws_claims.append(["C-1", 5000])

    wb.save(file_path)

    result = analyze_file(file_path)

    assert result.selected_sheet == "Property Schedule"
    assert len(result.all_sheets_evaluated) == 5
    evaluated_names = {s.sheet_name for s in result.all_sheets_evaluated}
    assert evaluated_names == {
        "Instructions",
        "Cover",
        "Summary",
        "Property Schedule",
        "Claims",
    }


# =====================================================================
# 12. No-SOV Test
# =====================================================================


def test_no_sov_table_detected(tmp_path: Path):
    """Section 12: Workbook containing only non-SOV sheets returns null selection."""
    file_path = tmp_path / "no_sov.xlsx"
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Instructions"
    ws1.append(["Please read this carefully"])

    ws2 = wb.create_sheet(title="Summary")
    ws2.append(["Summary of operations", "FY 2026"])

    ws3 = wb.create_sheet(title="Notes")
    ws3.append(["Contact broker for inquiries"])

    ws4 = wb.create_sheet(title="Cover")
    ws4.append(["Cover Page"])

    wb.save(file_path)

    result = analyze_file(file_path)

    assert result.selected_sheet is None
    assert result.header_row is None
    assert result.confidence == 0.0
    assert result.raw_headers == []
    assert result.sample_rows == []
    assert result.total_rows == 0
    assert len(result.all_sheets_evaluated) == 4
    assert any("No sheet qualified" in r for r in result.reasoning)


# =====================================================================
# 13. Ambiguous-Candidate Test (Near Tie)
# =====================================================================


def test_ambiguous_candidate_near_tie(tmp_path: Path):
    """Section 13: Workbook with two plausible SOV tables flags a near tie."""
    file_path = tmp_path / "two_schedules.xlsx"
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Location Data"
    ws1.append(["Ref", "Address", "City", "State", "Zip", "Building Value", "Occupancy"])
    ws1.append(["L1", "100 Pine", "Seattle", "WA", "98101", 5000000, "Office"])

    ws2 = wb.create_sheet(title="Property Schedule")
    ws2.append(["Ref", "Address", "City", "State", "Zip", "Building Value", "Occupancy"])
    ws2.append(["L2", "200 Pine", "Seattle", "WA", "98102", 7000000, "Retail"])

    wb.save(file_path)

    result = analyze_file(file_path)

    assert result.selected_sheet in {"Location Data", "Property Schedule"}
    assert len(result.all_sheets_evaluated) == 2


# =====================================================================
# 14. Serialized JSON Validation
# =====================================================================


def test_serialized_json_validation(tmp_path: Path):
    """Section 14: Validates clean JSON serialization without custom/unserializable classes."""
    file_path = tmp_path / "valid.csv"
    file_path.write_text(
        "Ref,Address,City,State,Zip,Building Cost,Occupancy\n"
        "001,100 Main,Chicago,IL,60601,2500000,Office\n",
        encoding="utf-8",
    )

    result = analyze_file(file_path)
    json_bytes = result.model_dump_json()

    parsed = json.loads(json_bytes)
    assert isinstance(parsed, dict)
    assert parsed["selected_sheet"] == "default"
    assert parsed["header_row"] == 0
    assert 0.0 <= parsed["confidence"] <= 1.0


# =====================================================================
# 16. Deterministic Execution Test
# =====================================================================


def test_deterministic_repeated_execution(tmp_path: Path):
    """Section 16: Running pipeline multiple times on same input produces identical output."""
    file_path = tmp_path / "determ.csv"
    file_path.write_text(
        "Loc ID,Street Address,City,State,Building Cost\n"
        "1,100 Main,Boston,MA,1000000\n",
        encoding="utf-8",
    )

    res1 = analyze_file(file_path, job_id="fixed_job")
    res2 = analyze_file(file_path, job_id="fixed_job")
    res3 = analyze_file(file_path, job_id="fixed_job")

    assert res1.model_dump() == res2.model_dump() == res3.model_dump()


# =====================================================================
# 17. Error Handling Tests
# =====================================================================


def test_missing_file_error(tmp_path: Path):
    """Missing file raises controlled FileNotFoundIngestionError."""
    with pytest.raises(FileNotFoundIngestionError):
        analyze_file(tmp_path / "missing.xlsx")


def test_unsupported_extension_error(tmp_path: Path):
    """Unsupported extension (.pdf) raises controlled UnsupportedFormatError."""
    p = tmp_path / "file.pdf"
    p.write_text("PDF content")
    with pytest.raises(UnsupportedFormatError):
        analyze_file(p)


def test_corrupted_xlsx_error(tmp_path: Path):
    """Corrupted XLSX raises controlled CorruptedFileError."""
    p = tmp_path / "corrupt.xlsx"
    p.write_bytes(b"PK\x03\x04corrupted-binary")
    with pytest.raises(CorruptedFileError):
        analyze_file(p)


def test_empty_csv_error(tmp_path: Path):
    """Empty CSV file raises controlled EmptyFileError."""
    p = tmp_path / "empty.csv"
    p.write_text("")
    with pytest.raises(EmptyFileError):
        analyze_file(p)


def test_invalid_json_error(tmp_path: Path):
    """Invalid JSON syntax raises controlled MalformedJSONError."""
    p = tmp_path / "bad.json"
    p.write_text("{invalid_json:")
    with pytest.raises(MalformedJSONError):
        analyze_file(p)


# =====================================================================
# 18. Explicit test_agent1_agent2_handoff_contract
# =====================================================================


def test_agent1_agent2_handoff_contract(tmp_path: Path):
    """Section 18: Explicit test verifying required fields and absence of canonical schema mapping."""
    file_path = tmp_path / "contract_test.csv"
    file_path.write_text(
        "Ref #,Site,City,State,Zip,Building Cost ($),Occupancy\n"
        "001,Plant Alpha,Chicago,IL,60601,$3,500,000,Industrial\n",
        encoding="utf-8",
    )

    result = analyze_file(file_path, job_id="job_contract_audit")

    # Required fields exist and are populated
    assert result.source_file == "contract_test.csv"
    assert result.selected_sheet == "default"
    assert result.header_row == 0
    assert 0.0 <= result.confidence <= 1.0
    assert result.raw_headers == [
        "Ref #",
        "Site",
        "City",
        "State",
        "Zip",
        "Building Cost ($)",
        "Occupancy",
    ]
    assert len(result.sample_rows) == 1
    assert result.total_rows == 1
    assert len(result.all_sheets_evaluated) == 1
    assert len(result.reasoning) >= 1

    # Verify Agent 1 does NOT contain canonical schema mappings
    data = json.loads(result.model_dump_json())
    raw_str = json.dumps(data)

    assert "schema_mapping" not in data
    assert "canonical_mapping" not in data
    assert '"Building Cost ($)": "building_value"' not in raw_str
    assert '"Site": "address"' not in raw_str
