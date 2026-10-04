"""
tests/test_excel_output.py
--------------------------
Unit tests for ExcelOutputService, ExcelService, and STANDARD_FIELDS.
Ported from Member 3 standalone test_excel_output.py.
"""

import copy
from pathlib import Path
import openpyxl
import pandas as pd
import pytest

from app.services.excel_service import (
    STANDARD_FIELDS,
    STRING_FIELDS,
    WORKSHEET_NAME,
    ExcelOutputService,
    ExcelService,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

@pytest.fixture
def service() -> ExcelOutputService:
    return ExcelOutputService()


@pytest.fixture
def insurance_record() -> dict:
    """A realistic single SOV record with 7 of the 17 fields populated."""
    return {
        "Reference":             "LOC-001",
        "Address":               "123 Main St",
        "City":                  "Boston",
        "State":                 "MA",
        "Zip":                   "02110",
        "Building Value":        1_250_000.0,
        "Fire Sprinklers (Y/N)": "Y",
    }


@pytest.fixture
def full_record() -> dict:
    """A record with all 17 standard fields populated."""
    return {
        "Reference":          "LOC-002",
        "Address":            "456 Oak Ave",
        "City":               "Springfield",
        "State":              "IL",
        "Zip":                "62701",
        "County":             "Sangamon",
        "Country":            "USA",
        "Building Value":     900_000.0,
        "Contents":           150_000.0,
        "BI":                 50_000.0,
        "Occupancy":          "Office",
        "Construction":       "Frame",
        "Storeys":            3,
        "Number of Buildings": 1,
        "Year Built":         1995,
        "Fire Sprinklers (Y/N)": "N",
        "Other":              "",
    }


# ===========================================================================
# 1. STANDARD_FIELDS schema
# ===========================================================================

class TestStandardFields:

    def test_standard_fields_contains_exactly_17(self):
        """STANDARD_FIELDS must have exactly 17 entries."""
        assert len(STANDARD_FIELDS) == 17

    def test_standard_fields_exact_names(self):
        """Every field name must match the spec exactly (case-sensitive)."""
        expected = [
            "Reference", "Address", "City", "State", "Zip",
            "County", "Country", "Building Value", "Contents", "BI",
            "Occupancy", "Construction", "Storeys", "Number of Buildings",
            "Year Built", "Fire Sprinklers (Y/N)", "Other",
        ]
        assert STANDARD_FIELDS == expected

    def test_standard_fields_exact_order(self):
        """Column order must exactly match the specification."""
        assert STANDARD_FIELDS[0]  == "Reference"
        assert STANDARD_FIELDS[7]  == "Building Value"
        assert STANDARD_FIELDS[15] == "Fire Sprinklers (Y/N)"
        assert STANDARD_FIELDS[16] == "Other"

    def test_standard_fields_no_duplicates(self):
        """STANDARD_FIELDS must not contain duplicate column names."""
        assert len(STANDARD_FIELDS) == len(set(STANDARD_FIELDS))


# ===========================================================================
# 2. Output directory and file creation
# ===========================================================================

class TestFileCreation:

    def test_creates_output_directory_if_missing(self, service, tmp_path):
        """write() must create the output directory automatically."""
        nested = tmp_path / "deep" / "nested" / "dir"
        output = nested / "Cleaned_SOV.xlsx"
        assert not nested.exists()

        service.write(records=[], output_path=output)

        assert nested.exists()
        assert output.exists()

    def test_creates_xlsx_file(self, service, tmp_path):
        """write() must produce a file with a .xlsx extension."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        assert output.exists()
        assert output.suffix == ".xlsx"

    def test_file_is_valid_xlsx(self, service, tmp_path):
        """The produced file must be openable as an Excel workbook."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        wb = openpyxl.load_workbook(output)
        assert wb is not None

    def test_returns_resolved_path(self, service, tmp_path):
        """write() must return a resolved (absolute) Path object."""
        output = tmp_path / "out.xlsx"
        returned = service.write(records=[], output_path=output)
        assert returned.is_absolute()


# ===========================================================================
# 3. Worksheet name
# ===========================================================================

class TestWorksheetName:

    def test_worksheet_name_constant(self):
        """WORKSHEET_NAME must be exactly 'Cleaned SOV'."""
        assert WORKSHEET_NAME == "Cleaned SOV"

    def test_worksheet_is_named_correctly(self, service, tmp_path):
        """The workbook must contain a sheet named 'Cleaned SOV'."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        wb = openpyxl.load_workbook(output)
        assert "Cleaned SOV" in wb.sheetnames

    def test_only_one_worksheet(self, service, tmp_path):
        """The workbook must contain exactly one sheet."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        wb = openpyxl.load_workbook(output)
        assert len(wb.sheetnames) == 1


# ===========================================================================
# 4. Header row
# ===========================================================================

class TestHeaderRow:

    def _headers(self, path) -> list:
        """Read the header row from the 'Cleaned SOV' sheet via openpyxl."""
        wb = openpyxl.load_workbook(path)
        ws = wb["Cleaned SOV"]
        return [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    def test_header_contains_exactly_17_columns(self, service, tmp_path):
        """The header row must have exactly 17 cells."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        assert len(self._headers(output)) == 17

    def test_header_order_is_correct(self, service, tmp_path):
        """Header column order must exactly match STANDARD_FIELDS."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        assert self._headers(output) == STANDARD_FIELDS

    def test_header_does_not_contain_audit_columns(self, service, tmp_path):
        """Audit / internal fields must NOT appear in the header."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        forbidden = {
            "source_field", "target_field", "confidence", "reasoning",
            "decision", "approver", "timestamp", "approval_status",
            "before_value", "after_value", "transformation",
        }
        headers_lower = {h.lower() for h in self._headers(output) if h}
        assert headers_lower.isdisjoint({f.lower() for f in forbidden})


# ===========================================================================
# 5. Data rows
# ===========================================================================

class TestDataRows:

    def _read(self, path) -> pd.DataFrame:
        return pd.read_excel(path, sheet_name="Cleaned SOV", engine="openpyxl")

    def test_empty_input_creates_zero_data_rows(self, service, tmp_path):
        """Empty records list → workbook with header only (0 data rows)."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        df = self._read(output)
        assert len(df) == 0

    def test_empty_input_still_has_17_columns(self, service, tmp_path):
        """Empty records list → workbook still has exactly 17 columns."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[], output_path=output)
        df = self._read(output)
        assert list(df.columns) == STANDARD_FIELDS

    def test_single_record_produces_one_data_row(self, service, tmp_path, insurance_record):
        """One record → exactly one data row in the output."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[insurance_record], output_path=output)
        df = self._read(output)
        assert len(df) == 1

    def test_multiple_records_produce_correct_row_count(self, service, tmp_path, full_record):
        """Three records → exactly three data rows."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        records = [full_record, full_record, full_record]
        service.write(records=records, output_path=output)
        df = self._read(output)
        assert len(df) == 3


# ===========================================================================
# 6. Missing fields → blank cells
# ===========================================================================

class TestMissingFields:

    def _read(self, path) -> pd.DataFrame:
        return pd.read_excel(path, sheet_name="Cleaned SOV", engine="openpyxl")

    def test_missing_fields_become_blank(self, service, tmp_path):
        """A record with only 2 fields → the other 15 must be blank (NaN)."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        partial = {"Reference": "LOC-001", "Building Value": 1_000_000}
        service.write(records=[partial], output_path=output)
        df = self._read(output)

        assert list(df.columns) == STANDARD_FIELDS
        for col in STANDARD_FIELDS:
            if col not in ("Reference", "Building Value"):
                assert pd.isna(df.loc[0, col]), f"Expected blank for {col!r}"

    def test_all_17_columns_present_even_for_partial_record(self, service, tmp_path):
        """Even a minimal record must result in 17-column output."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[{"Reference": "LOC-X"}], output_path=output)
        df = self._read(output)
        assert list(df.columns) == STANDARD_FIELDS


# ===========================================================================
# 7. Extra / internal fields are excluded
# ===========================================================================

class TestExtraFieldsExcluded:

    def _read(self, path) -> pd.DataFrame:
        return pd.read_excel(path, sheet_name="Cleaned SOV", engine="openpyxl")

    def test_extra_fields_not_in_output_columns(self, service, tmp_path):
        """Columns not in STANDARD_FIELDS must be silently dropped."""
        record = {
            "Reference": "LOC-001",
            "Building Value": 1_000_000,
            "confidence": 0.95,
            "reasoning":  "alias match",
            "approver":   "Member 3",
            "timestamp":  "2026-01-01",
        }
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[record], output_path=output)
        df = self._read(output)
        assert list(df.columns) == STANDARD_FIELDS
        assert "confidence" not in df.columns
        assert "reasoning"  not in df.columns
        assert "approver"   not in df.columns

    def test_output_has_exactly_17_columns_with_extra_input(self, service, tmp_path):
        """Column count must be exactly 17 even when input has extra keys."""
        record = {f"extra_{i}": i for i in range(20)}
        record["Reference"] = "LOC-999"
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[record], output_path=output)
        df = self._read(output)
        assert len(df.columns) == 17


# ===========================================================================
# 8. Source data safety
# ===========================================================================

class TestSourceDataSafety:

    def test_write_does_not_mutate_source_records(self, service, tmp_path, insurance_record):
        """write() must not change the caller's record dicts."""
        snapshot = copy.deepcopy(insurance_record)
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[insurance_record], output_path=output)
        assert insurance_record == snapshot

    def test_write_does_not_mutate_source_list(self, service, tmp_path, full_record):
        """The records list itself must not be mutated."""
        records = [copy.deepcopy(full_record), copy.deepcopy(full_record)]
        original_length = len(records)
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=records, output_path=output)
        assert len(records) == original_length


# ===========================================================================
# 9. Value integrity
# ===========================================================================

class TestValueIntegrity:

    def _read(self, path) -> pd.DataFrame:
        return pd.read_excel(path, sheet_name="Cleaned SOV", engine="openpyxl")

    def test_numeric_value_is_preserved(self, service, tmp_path):
        """Numeric Building Value must survive the round-trip to Excel."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[{"Building Value": 1_250_000.0}], output_path=output)
        df = self._read(output)
        assert df.loc[0, "Building Value"] == 1_250_000.0

    def test_string_value_is_preserved(self, service, tmp_path):
        """String Reference must survive the round-trip to Excel."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[{"Reference": "LOC-ABC"}], output_path=output)
        df = self._read(output)
        assert df.loc[0, "Reference"] == "LOC-ABC"

    def test_output_readable_by_pandas(self, service, tmp_path, insurance_record):
        """pd.read_excel on the output must succeed without error."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[insurance_record], output_path=output)
        df = pd.read_excel(output, sheet_name="Cleaned SOV", engine="openpyxl")
        assert isinstance(df, pd.DataFrame)


# ===========================================================================
# 10. Insurance-like end-to-end example
# ===========================================================================

class TestInsuranceExample:

    def _read(self, path) -> pd.DataFrame:
        dtype_map = {col: str for col in STRING_FIELDS}
        return pd.read_excel(
            path, sheet_name="Cleaned SOV", engine="openpyxl", dtype=dtype_map
        )

    def test_insurance_record_fields_in_output(self, service, tmp_path, insurance_record):
        """All 7 populated fields must appear with correct values."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[insurance_record], output_path=output)
        df = self._read(output)

        assert df.loc[0, "Reference"]             == "LOC-001"
        assert df.loc[0, "Address"]               == "123 Main St"
        assert df.loc[0, "City"]                  == "Boston"
        assert df.loc[0, "State"]                 == "MA"
        assert str(df.loc[0, "Zip"])              == "02110"
        assert df.loc[0, "Building Value"]        == 1_250_000.0
        assert df.loc[0, "Fire Sprinklers (Y/N)"] == "Y"

    def test_insurance_record_unpopulated_fields_are_blank(self, service, tmp_path, insurance_record):
        """Fields not in the insurance record must be blank (NaN)."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[insurance_record], output_path=output)
        df = self._read(output)

        populated = set(insurance_record.keys())
        for col in STANDARD_FIELDS:
            if col not in populated:
                val = df.loc[0, col]
                is_blank = pd.isna(val) or str(val).strip().lower() == "nan"
                assert is_blank, f"Expected blank for {col!r}, got {val!r}"

    def test_insurance_output_has_17_columns(self, service, tmp_path, insurance_record):
        """Insurance example output must have exactly 17 columns."""
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=[insurance_record], output_path=output)
        df = self._read(output)
        assert len(df.columns) == 17

    def test_two_location_records(self, service, tmp_path):
        """Multiple locations must all appear as separate rows."""
        records = [
            {"Reference": "LOC-001", "Building Value": 1_250_000.0, "State": "MA"},
            {"Reference": "LOC-002", "Building Value":   900_000.0, "State": "TX"},
        ]
        output = tmp_path / "Cleaned_SOV.xlsx"
        service.write(records=records, output_path=output)
        df = self._read(output)

        assert len(df) == 2
        assert df.loc[0, "Reference"] == "LOC-001"
        assert df.loc[1, "Reference"] == "LOC-002"
        assert df.loc[0, "Building Value"] == 1_250_000.0
        assert df.loc[1, "Building Value"] ==   900_000.0


# ===========================================================================
# 11. Team ExcelService specific tests
# ===========================================================================

class TestTeamExcelService:

    def test_excel_service_sheet_methods(self):
        service = ExcelService()
        sheets = service.get_sheet_names("dummy.xlsx")
        assert len(sheets) >= 1
        sample = service.read_sheet_sample("dummy.xlsx", "Sheet1")
        assert isinstance(sample, list)

    def test_excel_service_export_cleaned_sov(self, tmp_path, insurance_record):
        service = ExcelService()
        output_file = tmp_path / "exported" / "Cleaned_SOV.xlsx"
        ret = service.export_cleaned_sov([insurance_record], output_file)
        assert str(ret) == str(output_file.resolve())
        assert output_file.exists()

        wb = openpyxl.load_workbook(output_file)
        assert "Cleaned SOV" in wb.sheetnames
        df = pd.read_excel(output_file, sheet_name="Cleaned SOV")
        assert list(df.columns) == STANDARD_FIELDS
        assert len(df) == 1
        assert df.loc[0, "Reference"] == "LOC-001"
