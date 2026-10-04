"""Focused tests for Agent 1 Universal Ingestion Layer.

Validates XLSX, CSV, and JSON ingestion without testing downstream
header detection or semantic mapping.
"""

from pathlib import Path
import openpyxl
import pytest

from app.agents.sheet_intelligence.errors import (
    CorruptedFileError,
    EmptyFileError,
    FileNotFoundIngestionError,
    MalformedJSONError,
    UnsupportedFormatError,
    UnsupportedJSONStructureError,
)
from app.agents.sheet_intelligence.ingestion import UniversalIngestor


@pytest.fixture
def ingestor() -> UniversalIngestor:
    return UniversalIngestor()


# =====================================================================
# 1 & 2. Valid XLSX: Single Sheet & Multiple Sheets
# =====================================================================


def test_valid_xlsx_single_sheet(tmp_path: Path, ingestor: UniversalIngestor):
    """Test 1: Valid XLSX with one sheet."""
    file_path = tmp_path / "single_sheet.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Property Schedule"
    ws.append(["Reference", "Site", "Building Cost"])
    ws.append(["LOC-001", "100 Main St", 1500000])
    ws.append(["LOC-002", "200 Park Ave", 2500000])
    wb.save(file_path)

    result = ingestor.ingest(file_path)
    assert result.file_type == "xlsx"
    assert "Property Schedule" in result.tables
    grid = result.tables["Property Schedule"]
    assert grid.total_rows == 3
    assert grid.total_cols == 3
    assert grid.matrix[0] == ["Reference", "Site", "Building Cost"]
    assert grid.matrix[1] == ["LOC-001", "100 Main St", 1500000]


def test_valid_xlsx_multiple_sheets(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test 2 & 17: Valid XLSX with multiple sheets and preservation of sheet names."""
    file_path = tmp_path / "multi_sheet.xlsx"
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Instructions"
    ws1.append(["Please fill all property fields below."])

    ws2 = wb.create_sheet(title="Locations")
    ws2.append(["Ref", "Address", "TIV"])
    ws2.append(["L-1", "Market St", 5000000])

    ws3 = wb.create_sheet(title="Claims History")
    ws3.append(["Claim ID", "Loss Amount"])
    ws3.append(["C-99", 12000])

    wb.save(file_path)

    result = ingestor.ingest(file_path)
    assert set(result.tables.keys()) == {
        "Instructions",
        "Locations",
        "Claims History",
    }
    assert result.tables["Locations"].matrix[0] == ["Ref", "Address", "TIV"]
    assert result.tables["Claims History"].matrix[1] == ["C-99", 12000]


# =====================================================================
# 3 & 4. XLSX: Metadata Rows and Blank Rows Preservation
# =====================================================================


def test_xlsx_with_metadata_and_blank_rows(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test 3, 4, 16: Metadata rows before actual data, blank rows preserved in exact row positions."""
    file_path = tmp_path / "metadata_sov.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"

    ws.append(["Broker Name: Risk Solutions LLC"])  # Row 0
    ws.append(["Property Schedule Submission"])  # Row 1
    ws.append(["Date: 2026-10-01"])  # Row 2
    ws.append([])  # Row 3 (blank row)
    ws.append(
        ["Reference", "Site", "Location", "Building Cost"]
    )  # Row 4 (headers)
    ws.append(["P001", "HQ", "100 Main St", 10000000])  # Row 5 (record)

    wb.save(file_path)

    result = ingestor.ingest(file_path)
    grid = result.tables["Schedule"]

    assert grid.total_rows == 6
    # Row 0
    assert grid.matrix[0][0] == "Broker Name: Risk Solutions LLC"
    # Row 1
    assert grid.matrix[1][0] == "Property Schedule Submission"
    # Row 2
    assert grid.matrix[2][0] == "Date: 2026-10-01"
    # Row 3 must be preserved as empty/None cells, NOT discarded!
    assert all(cell is None for cell in grid.matrix[3])
    # Row 4 must remain at index 4!
    assert grid.matrix[4][:4] == [
        "Reference",
        "Site",
        "Location",
        "Building Cost",
    ]
    # Row 5
    assert grid.matrix[5][:4] == ["P001", "HQ", "100 Main St", 10000000]


# =====================================================================
# 5 & 6. CSV: Valid CSV and CSV with Metadata Rows
# =====================================================================


def test_valid_csv(tmp_path: Path, ingestor: UniversalIngestor):
    """Test 5: Valid CSV ingested into one logical table."""
    file_path = tmp_path / "valid.csv"
    csv_content = (
        "Reference,Location,Building Cost,Contents,BI\n"
        "LOC-101,1200 Market St,4500000,500000,250000\n"
        "LOC-102,400 Pine St,2100000,250000,100000\n"
    )
    file_path.write_text(csv_content, encoding="utf-8")

    result = ingestor.ingest(file_path)
    assert result.file_type == "csv"
    assert "default" in result.tables
    grid = result.tables["default"]
    assert grid.total_rows == 3
    assert grid.matrix[0] == [
        "Reference",
        "Location",
        "Building Cost",
        "Contents",
        "BI",
    ]
    assert grid.matrix[1] == [
        "LOC-101",
        "1200 Market St",
        "4500000",
        "500000",
        "250000",
    ]


def test_csv_with_metadata_rows_and_blanks(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test 6 & 16: CSV with metadata rows and blank lines before the table header."""
    file_path = tmp_path / "metadata.csv"
    csv_content = (
        "Acme Brokerage Corp\n"
        "Commercial Property Schedule\n"
        "\n"
        "Reference,Site,Building Cost\n"
        "R-1,Warehouse A,8000000\n"
    )
    file_path.write_text(csv_content, encoding="utf-8")

    result = ingestor.ingest(file_path)
    grid = result.tables["default"]

    assert grid.total_rows == 5
    assert grid.matrix[0][0] == "Acme Brokerage Corp"
    assert grid.matrix[1][0] == "Commercial Property Schedule"
    # Row 2 is blank line
    assert all(c is None for c in grid.matrix[2])
    # Row 3 is headers
    assert grid.matrix[3][:3] == ["Reference", "Site", "Building Cost"]
    # Row 4 is record
    assert grid.matrix[4][:3] == ["R-1", "Warehouse A", "8000000"]


def test_csv_semicolon_and_bom_encoding(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test CSV with semicolon delimiter and UTF-8 BOM encoding."""
    file_path = tmp_path / "semicolon_bom.csv"
    content = (
        "Reference;Address;Building Value\n"
        "LOC-1;100 Main St;5000000\n"
        "LOC-2;200 Oak St;3500000\n"
    )
    file_path.write_text(content, encoding="utf-8-sig")

    result = ingestor.ingest(file_path)
    grid = result.tables["default"]
    assert grid.total_rows == 3
    assert grid.matrix[0] == ["Reference", "Address", "Building Value"]
    assert grid.matrix[1] == ["LOC-1", "100 Main St", "5000000"]


# =====================================================================
# 7 & 8. JSON: List of Records & Supported Dictionary Structures
# =====================================================================


def test_valid_json_list_of_records(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test 7: Valid JSON list of records."""
    file_path = tmp_path / "records.json"
    content = """[
        {"Reference": "P001", "Address": "100 Main St", "Building Cost": 1500000},
        {"Reference": "P002", "Address": "200 Park St", "Building Cost": 2000000}
    ]"""
    file_path.write_text(content, encoding="utf-8")

    result = ingestor.ingest(file_path)
    assert result.file_type == "json"
    grid = result.tables["default"]
    assert grid.total_rows == 3
    # Row 0: Original key names
    assert grid.matrix[0] == ["Reference", "Address", "Building Cost"]
    # Row 1 and 2: Record values
    assert grid.matrix[1] == ["P001", "100 Main St", 1500000]
    assert grid.matrix[2] == ["P002", "200 Park St", 2000000]


def test_valid_json_dictionary_with_tables(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test 8: Supported dictionary containing named table lists."""
    file_path = tmp_path / "nested_tables.json"
    content = """{
        "Property Schedule": [
            {"Site": "Facility Alpha", "Building Cost": 12000000, "Occupancy": "Industrial"},
            {"Site": "Facility Beta", "Building Cost": 4500000, "Occupancy": "Office"}
        ],
        "Claims": [
            {"Claim ID": "CLM-01", "Amount": 50000}
        ]
    }"""
    file_path.write_text(content, encoding="utf-8")

    result = ingestor.ingest(file_path)
    assert "Property Schedule" in result.tables
    assert "Claims" in result.tables

    sov_grid = result.tables["Property Schedule"]
    assert sov_grid.matrix[0] == ["Site", "Building Cost", "Occupancy"]
    assert sov_grid.matrix[1] == ["Facility Alpha", 12000000, "Industrial"]


def test_valid_json_2d_array_and_column_oriented(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test 2D array and column-oriented JSON shapes."""
    # 2D array
    file_path_2d = tmp_path / "array2d.json"
    file_path_2d.write_text(
        '[["Ref", "Val"], ["001", 100], ["002", 200]]', encoding="utf-8"
    )
    res_2d = ingestor.ingest(file_path_2d)
    assert res_2d.tables["default"].matrix[0] == ["Ref", "Val"]
    assert res_2d.tables["default"].matrix[1] == ["001", 100]

    # Column-oriented
    file_path_col = tmp_path / "column_oriented.json"
    file_path_col.write_text(
        '{"Ref": ["R1", "R2"], "Cost": [5000, 7000]}', encoding="utf-8"
    )
    res_col = ingestor.ingest(file_path_col)
    assert res_col.tables["default"].matrix[0] == ["Ref", "Cost"]
    assert res_col.tables["default"].matrix[1] == ["R1", 5000]
    assert res_col.tables["default"].matrix[2] == ["R2", 7000]


# =====================================================================
# 9, 10, 11. Empty Files (XLSX, CSV, JSON)
# =====================================================================


def test_empty_xlsx_sheet(tmp_path: Path, ingestor: UniversalIngestor):
    """Test 9: Empty XLSX sheet returns empty matrix without fabricating rows."""
    file_path = tmp_path / "empty_sheet.xlsx"
    wb = openpyxl.Workbook()
    # active sheet has 0 cells populated
    wb.save(file_path)

    result = ingestor.ingest(file_path)
    assert "Sheet" in result.tables
    assert result.tables["Sheet"].matrix == []
    assert result.tables["Sheet"].total_rows == 0


def test_empty_csv_zero_bytes(tmp_path: Path, ingestor: UniversalIngestor):
    """Test 10: Empty CSV file (0 bytes) raises controlled EmptyFileError."""
    file_path = tmp_path / "empty.csv"
    file_path.write_text("", encoding="utf-8")

    with pytest.raises(EmptyFileError) as exc_info:
        ingestor.ingest(file_path)
    assert "0 bytes" in str(exc_info.value)


def test_empty_json(tmp_path: Path, ingestor: UniversalIngestor):
    """Test 11: Empty JSON file (0 bytes) or empty array."""
    # 0 bytes
    file_path_empty = tmp_path / "empty.json"
    file_path_empty.write_text("", encoding="utf-8")
    with pytest.raises(EmptyFileError):
        ingestor.ingest(file_path_empty)

    # Empty array []
    file_path_array = tmp_path / "empty_array.json"
    file_path_array.write_text("[]", encoding="utf-8")
    res = ingestor.ingest(file_path_array)
    assert res.tables["default"].matrix == []


# =====================================================================
# 12, 13, 14. Malformed Input, Missing File, Unsupported Extension
# =====================================================================


def test_invalid_json_syntax(tmp_path: Path, ingestor: UniversalIngestor):
    """Test 12: Invalid JSON syntax raises controlled MalformedJSONError."""
    file_path = tmp_path / "broken.json"
    file_path.write_text("{'key': incomplete json...", encoding="utf-8")

    with pytest.raises(MalformedJSONError) as exc_info:
        ingestor.ingest(file_path)
    assert "Invalid JSON syntax" in str(exc_info.value)


def test_unsupported_json_structure(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test unsupported arbitrary/non-tabular JSON structures."""
    file_path = tmp_path / "arbitrary.json"
    # A single number or primitive is not a table
    file_path.write_text("12345", encoding="utf-8")

    with pytest.raises(UnsupportedJSONStructureError):
        ingestor.ingest(file_path)


def test_missing_file(tmp_path: Path, ingestor: UniversalIngestor):
    """Test 13: Missing file raises controlled FileNotFoundIngestionError."""
    non_existent = tmp_path / "does_not_exist.xlsx"
    with pytest.raises(FileNotFoundIngestionError):
        ingestor.ingest(non_existent)


def test_unsupported_extension(tmp_path: Path, ingestor: UniversalIngestor):
    """Test 14: Unsupported extension (.pdf, .parquet, .txt) raises UnsupportedFormatError."""
    file_path = tmp_path / "document.pdf"
    file_path.write_text("PDF content", encoding="utf-8")

    with pytest.raises(UnsupportedFormatError) as exc_info:
        ingestor.ingest(file_path)
    assert "Unsupported file format '.pdf'" in str(exc_info.value)


def test_corrupted_xlsx(tmp_path: Path, ingestor: UniversalIngestor):
    """Corrupted XLSX binary raises controlled CorruptedFileError."""
    file_path = tmp_path / "corrupted.xlsx"
    file_path.write_bytes(b"PK\x03\x04not-a-valid-zip-content-here")

    with pytest.raises(CorruptedFileError):
        ingestor.ingest(file_path)


# =====================================================================
# 15. Data Preservation: Verbatim Header & Value Preservation
# =====================================================================


def test_data_preservation_raw_headers_and_values(
    tmp_path: Path, ingestor: UniversalIngestor
):
    """Test 15: Agent 1 preserves exact raw text ('Building Cost', 'Site', 'BI')

    without renaming to canonical SOV fields ('Building Value', 'Address', etc.).
    """
    file_path = tmp_path / "raw_names.csv"
    file_path.write_text(
        "Building Cost,Site,BI,Occupancy Description\n"
        "1500000,Location A,250000,Manufacturing\n",
        encoding="utf-8",
    )

    result = ingestor.ingest(file_path)
    headers = result.tables["default"].matrix[0]

    # Verify no canonical normalization was done
    assert "Building Cost" in headers
    assert "Building Value" not in headers

    assert "Site" in headers
    assert "Address" not in headers

    assert "BI" in headers
    assert "Business Interruption" not in headers


# =====================================================================
# Canonical Dict Conversion Helper
# =====================================================================


def test_to_dict_conversion(tmp_path: Path, ingestor: UniversalIngestor):
    """Verify that IngestedWorkbook.to_dict() matches the conceptual specification:

    {
        "sheet_name": [
            [cell1, cell2, ...],
            ...
        ]
    }
    """
    file_path = tmp_path / "simple.csv"
    file_path.write_text("A,B\n1,2\n3,4\n", encoding="utf-8")

    result = ingestor.ingest(file_path)
    raw_dict = result.to_dict()

    assert "default" in raw_dict
    assert isinstance(raw_dict["default"], list)
    assert raw_dict["default"] == [["A", "B"], ["1", "2"], ["3", "4"]]
