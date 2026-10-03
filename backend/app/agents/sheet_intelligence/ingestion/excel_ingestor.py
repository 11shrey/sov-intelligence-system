"""XLSX Ingestion implementation using openpyxl."""

import os
from pathlib import Path
from typing import Any
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.agents.sheet_intelligence.errors import (
    CorruptedFileError,
    EmptyFileError,
    FileNotFoundIngestionError,
)
from app.agents.sheet_intelligence.ingestion.base import (
    IngestedWorkbook,
    RawGrid,
    TableIngestor,
)


class ExcelIngestor(TableIngestor):
    """Ingests Excel (.xlsx) workbooks without assuming headers or mutating source content."""

    def ingest(
        self, file_path: str | Path, max_rows: int | None = None
    ) -> IngestedWorkbook:
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundIngestionError(
                f"File not found: {path}", file_path=str(path)
            )

        if path.stat().st_size == 0:
            raise EmptyFileError(
                f"File is empty (0 bytes): {path}", file_path=str(path)
            )

        try:
            wb = load_workbook(
                filename=str(path), read_only=True, data_only=True
            )
        except (
            InvalidFileException,
            zipfile.BadZipFile,
            KeyError,
            ValueError,
        ) as exc:
            raise CorruptedFileError(
                f"Cannot parse XLSX file: {exc}",
                file_path=str(path),
                details={"error": str(exc)},
            ) from exc

        try:
            sheet_names = wb.sheetnames
            if not sheet_names:
                raise EmptyFileError(
                    f"XLSX workbook has no sheets: {path}", file_path=str(path)
                )

            tables: dict[str, RawGrid] = {}

            for sheet_name in sheet_names:
                ws = wb[sheet_name]
                raw_rows: list[list[Any]] = []
                row_idx = 0

                for row_tuple in ws.iter_rows(values_only=True):
                    raw_rows.append(list(row_tuple))
                    row_idx += 1
                    if max_rows is not None and row_idx >= max_rows:
                        break

                # Strip trailing completely empty rows (ghost rows from formatting)
                # while strictly preserving all leading and interior blank rows
                last_data_idx = -1
                for idx, r in enumerate(raw_rows):
                    if any(
                        cell is not None and str(cell).strip() != "" for cell in r
                    ):
                        last_data_idx = idx

                if last_data_idx >= 0:
                    trimmed_rows = raw_rows[: last_data_idx + 1]
                else:
                    trimmed_rows = []

                tables[sheet_name] = RawGrid.from_rows(
                    name=sheet_name, rows=trimmed_rows
                )

            return IngestedWorkbook(
                source_path=str(path),
                file_type="xlsx",
                tables=tables,
            )
        finally:
            wb.close()
